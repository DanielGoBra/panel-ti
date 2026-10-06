import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter

ACCENT = "5B5BD6"
HEADER_BG = "F1F1EF"


def build_xlsx(report):
    rep = report
    k = rep["kpis"]
    wb = Workbook()

    # ---------- Hoja 1: Resumen ----------
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False

    ws["B2"] = rep["title"]
    ws["B2"].font = Font(size=18, bold=True, color=ACCENT)
    ws["B3"] = f"Período: {rep['period']['start']} → {rep['period']['end']}   ·   Generado: {rep['generated_at']}"
    ws["B3"].font = Font(size=9, color="7A7A7A")

    kpis = [
        ("Tareas totales", k["total_tasks"]),
        ("Completado (%)", k["completion_pct"]),
        ("Hechas", k["done"]),
        ("En curso", k["in_progress"]),
        ("Bloqueadas", k["blocked"]),
        ("Vencidas", k["overdue"]),
        ("Proyectos", k["projects_tracked"]),
    ]
    ws["B5"] = "Indicador"
    ws["C5"] = "Valor"
    for cell in ("B5", "C5"):
        ws[cell].fill = PatternFill("solid", start_color=HEADER_BG)
        ws[cell].font = Font(bold=True, size=10)
    for i, (lbl, val) in enumerate(kpis, start=6):
        ws[f"B{i}"] = lbl
        ws[f"C{i}"] = val
        ws[f"C{i}"].font = Font(bold=True, color=ACCENT)

    # Datos por proyecto (para gráfico)
    ws["E5"] = "Proyecto"
    for cell in ("E5", "F5", "G5"):
        ws[cell].fill = PatternFill("solid", start_color=HEADER_BG)
        ws[cell].font = Font(bold=True, size=10)
    ws["F5"] = "Hechas"
    ws["G5"] = "Abiertas"
    for i, p in enumerate(rep["by_project"], start=6):
        ws[f"E{i}"] = p["name"]
        ws[f"F{i}"] = p["done"]
        ws[f"G{i}"] = p["open"]

    n = len(rep["by_project"])
    if n:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Avance por proyecto"
        chart.style = 10
        data = Reference(ws, min_col=6, min_row=5, max_col=7, max_row=5 + n)
        cats = Reference(ws, min_col=5, min_row=6, max_row=5 + n)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 16
        chart.height = 9
        ws.add_chart(chart, "E" + str(9 + n))

    # ---------- Hoja 2: Proyectos ----------
    ws2 = wb.create_sheet("Proyectos")
    headers = ["Proyecto", "Total", "Hechas", "Abiertas", "Bloqueadas", "% Avance"]
    for c, h in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = PatternFill("solid", start_color=HEADER_BG)
        cell.font = Font(bold=True)
    for r, p in enumerate(rep["by_project"], start=2):
        pct = round(100 * p["done"] / p["total"], 1) if p["total"] else 0
        ws2.cell(row=r, column=1, value=p["name"])
        ws2.cell(row=r, column=2, value=p["total"])
        ws2.cell(row=r, column=3, value=p["done"])
        ws2.cell(row=r, column=4, value=p["open"])
        ws2.cell(row=r, column=5, value=p["blocked"])
        ws2.cell(row=r, column=6, value=pct / 100)
        ws2.cell(row=r, column=6).number_format = "0.0%"

    # ---------- Hoja 3: Alertas ----------
    ws3 = wb.create_sheet("Alertas")
    ws3["A1"] = "Tipo"
    ws3["B1"] = "Tarea"
    ws3["C1"] = "Detalle"
    for c in "ABC":
        ws3[c + "1"].fill = PatternFill("solid", start_color=HEADER_BG)
        ws3[c + "1"].font = Font(bold=True)
    r = 2
    for t in rep["overdue_tasks"]:
        ws3.cell(row=r, column=1, value="VENCIDA")
        ws3.cell(row=r, column=2, value=t["title"])
        ws3.cell(row=r, column=3, value=f"Vencía {t['due_date']} · prioridad {t['priority']}")
        r += 1
    for t in rep["blocked_tasks"]:
        ws3.cell(row=r, column=1, value="BLOQUEADA")
        ws3.cell(row=r, column=2, value=t["title"])
        ws3.cell(row=r, column=3, value=t.get("assignee", ""))
        r += 1

    # ---------- Hoja 4: Prioridades ----------
    ws4 = wb.create_sheet("Prioridades")
    ws4["A1"] = "Prioridad"
    ws4["B1"] = "Tareas"
    for c in "AB":
        ws4[c + "1"].fill = PatternFill("solid", start_color=HEADER_BG)
        ws4[c + "1"].font = Font(bold=True)
    labels = {"low": "Baja", "medium": "Media", "high": "Alta", "critical": "Crítica"}
    for i, (key, lbl) in enumerate(labels.items(), start=2):
        ws4.cell(row=i, column=1, value=lbl)
        ws4.cell(row=i, column=2, value=rep["by_priority"].get(key, 0))
    pie = PieChart()
    pie.title = "Distribución por prioridad"
    pie.add_data(Reference(ws4, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pie.set_categories(Reference(ws4, min_col=1, min_row=2, max_row=5))
    pie.width = 12
    pie.height = 8
    ws4.add_chart(pie, "D2")

    # Anchos de columna razonables
    for sheet in (ws, ws2, ws3, ws4):
        for col in range(1, 8):
            sheet.column_dimensions[get_column_letter(col)].width = 22
    ws.column_dimensions["E"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
