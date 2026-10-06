import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import parse_xml

ACCENT = "5B5BD6"
DARK = "323232"
MUTED = "7A7A7A"


def _set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:fill="{color_hex}"/>')
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def build_docx(report):
    rep = report
    k = rep["kpis"]
    doc = Document()

    # Estilo base
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor.from_string(DARK)

    # Título
    title = doc.add_heading(rep["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Período: {rep['period']['start']} → {rep['period']['end']}    ·    Generado: {rep['generated_at']}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    # Resumen ejecutivo (párrafo automático)
    doc.add_heading("Resumen ejecutivo", level=1)
    estado = "saludable"
    if k["blocked"] > 0:
        estado = "con bloqueos que requieren atención"
    elif k["overdue"] > k["total_tasks"] * 0.3:
        estado = "con retrasos significativos en el cronograma"
    resumen = (
        f"Durante el período analizado se gestionaron {k['total_tasks']} tareas distribuidas en "
        f"{k['projects_tracked']} proyectos, alcanzando un {k['completion_pct']}% de finalización. "
        f"Actualmente hay {k['in_progress']} tareas en curso, {k['todo']} pendientes y {k['blocked']} bloqueadas. "
        f"El estado general de la cartera es {estado}."
    )
    doc.add_paragraph(resumen)

    # KPIs
    doc.add_heading("Indicadores clave (KPIs)", level=1)
    kpi_data = [
        ("Tareas totales", k["total_tasks"], "total"),
        ("Completado", f"{k['completion_pct']}%", "pct"),
        ("Hechas", k["done"], "done"),
        ("En curso", k["in_progress"], "progress"),
        ("Bloqueadas", k["blocked"], "blocked"),
        ("Vencidas", k["overdue"], "overdue"),
    ]
    table = doc.add_table(rows=2, cols=len(kpi_data))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (lbl, val, _t) in enumerate(kpi_data):
        cell_lbl = table.rows[0].cells[i]
        cell_val = table.rows[1].cells[i]
        _set_cell_bg(cell_lbl, "F1F1EF")
        cell_lbl.text = lbl
        cell_val.text = str(val)
        for para in cell_val.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor.from_string(ACCENT)
        for para in cell_lbl.paragraphs:
            for run in para.runs:
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor.from_string(MUTED)

    # Gráfico de barras por proyecto (dibujado con celdas)
    doc.add_heading("Avance por proyecto", level=1)
    t2 = doc.add_table(rows=len(rep["by_project"]) + 1, cols=4)
    headers = ["Proyecto", "Total", "Hechas", "Avance"]
    for i, h in enumerate(headers):
        c = t2.rows[0].cells[i]
        c.text = h
        _set_cell_bg(c, "F1F1EF")
    for r_i, p in enumerate(rep["by_project"], start=1):
        pct = round(100 * p["done"] / p["total"]) if p["total"] else 0
        t2.rows[r_i].cells[0].text = p["name"]
        t2.rows[r_i].cells[1].text = str(p["total"])
        t2.rows[r_i].cells[2].text = str(p["done"])
        bar = "█" * max(1, pct // 5)
        cell = t2.rows[r_i].cells[3]
        cell.text = f"{bar} {pct}%"

    # Distribución por prioridad
    doc.add_heading("Distribución por prioridad", level=1)
    pr = rep["by_priority"]
    labels = {"low": "Baja", "medium": "Media", "high": "Alta", "critical": "Crítica"}
    for key, lbl in labels.items():
        v = pr.get(key, 0)
        pct = round(100 * v / k["total_tasks"]) if k["total_tasks"] else 0
        doc.add_paragraph(f"{lbl}: {v}  ({'█' * max(1, pct // 4)} {pct}%)")

    # Alertas
    if rep["overdue_tasks"]:
        doc.add_heading("Tareas vencidas", level=1)
        for t in rep["overdue_tasks"]:
            doc.add_paragraph(f"⚠ {t['title']} — vencía {t['due_date']} (prioridad {t['priority']})", style="List Bullet")
    if rep["blocked_tasks"]:
        doc.add_heading("Bloqueos que requieren atención del comité", level=1)
        for t in rep["blocked_tasks"]:
            doc.add_paragraph(f"🚫 {t['title']}" + (f" — {t['assignee']}" if t["assignee"] else ""), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
